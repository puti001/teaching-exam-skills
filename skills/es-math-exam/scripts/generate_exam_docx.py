#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
generate_exam_docx.py - 國小數學段考 Word 文件生成腳本 (含幾何配圖自動插入與 OMML 原生數學公式轉換)
"""
import sys
import json
import re
from pathlib import Path

# 強制 Windows 輸出 UTF-8 避免編碼錯誤
if sys.platform.startswith('win'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except AttributeError:
        pass

try:
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
except ImportError:
    print("❌ python-docx 未安裝。請執行: pip install python-docx")
    sys.exit(1)

def clean_math_text(text):
    """清理非公式文字的簡單標記"""
    if not text:
        return ""
    text = text.replace('{=}', '=')
    text = text.replace('{<=}', '≤')
    text = text.replace('{>=}', '≥')
    text = text.replace('{!=}', '≠')
    text = text.replace('{*}', '×')
    return text

def add_math_runs(paragraph, text):
    """將包含 {math} 標記的文字解析並以 OMML 原生公式或精緻格式插入段落中"""
    if not text:
        return
        
    # 首先將 {frac(a,b)} 切片
    parts = re.split(r'(\{frac\([^,]+,[^)]+\)\})', text)
    
    for part in parts:
        if part.startswith('{frac(') and part.endswith(')}'):
            m = re.match(r'\{frac\(([^,]+),([^)]+)\)\}', part)
            if m:
                num = m.group(1)
                den = m.group(2)
                # Word OMML 原生分數 XML
                omml_xml = (
                    f'<m:oMath {nsdecls("m")} xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                    f'<m:f>'
                    f'<m:num><m:r><w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/></w:rPr><m:t>{num}</m:t></m:r></m:num>'
                    f'<m:den><m:r><w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/></w:rPr><m:t>{den}</m:t></m:r></m:den>'
                    f'</m:f>'
                    f'</m:oMath>'
                )
                try:
                    omml_el = parse_xml(omml_xml)
                    paragraph._element.append(omml_el)
                except Exception:
                    # 降級處理為普通斜線分數
                    run = paragraph.add_run(f"{num}/{den}")
                    run.font.name = '新細明體'
                    run.font.size = Pt(11)
            else:
                run = paragraph.add_run(part)
                run.font.name = '新細明體'
                run.font.size = Pt(11)
        else:
            # 將國小未知數 {x} 或 {y} 降級/轉換為空的「□」方框，避免 x, y 字母造成超出範圍的誤會
            subparts = re.split(r'(\{x\}|\{y\})', part)
            for subpart in subparts:
                if subpart in ('{x}', '{y}'):
                    run = paragraph.add_run("□")
                    run.font.name = '新細明體'
                    run.font.size = Pt(11)
                else:
                    cleaned = clean_math_text(subpart)
                    if cleaned:
                        run = paragraph.add_run(cleaned)
                        run.font.name = '新細明體'
                        run.font.size = Pt(11)

def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = '微軟正黑體'
    run.font.size = Pt(18)
    run.bold = True
    return p

def add_subtitle(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = '微軟正黑體'
    run.font.size = Pt(11)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = '微軟正黑體'
    run.font.size = Pt(12)
    run.bold = True
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    return p

def add_question(doc, text, num, points=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    pts_str = f"({points}分)" if points else ""
    p.add_run(f"{num}. ")
    add_math_runs(p, text)
    if pts_str:
        p.add_run(f" {pts_str}")
    return p

def insert_geometry_image(doc, fig_id, q, geom_dir):
    """如果幾何圖形 PNG 存在，將其插入 Word 中"""
    if not geom_dir:
        return
    png_path = geom_dir / f"{fig_id}.png"
    if png_path.exists():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p_img.add_run()
        run.add_picture(str(png_path), width=Cm(6.0))
        
        geo_data = q.get('geometry', {})
        caption = geo_data.get('caption', '')
        if caption:
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_cap = p_cap.add_run(caption)
            run_cap.font.name = '新細明體'
            run_cap.font.size = Pt(9)
            run_cap.font.italic = True
            p_cap.paragraph_format.space_after = Pt(6)
        print(f"🖼️  已自動插入幾何圖形: {png_path.name}")

def create_exam_paper(data, out_dir, geom_dir):
    """產生題目卷"""
    doc = Document()
    
    title_text = f"中華民國 114 學年度第 {data.get('semester', '1')} 學期第 {data.get('exam_number', '2')} 次段考"
    add_title(doc, f"{data.get('grade', '五年級')} 數學科 題目卷")
    add_subtitle(doc, f"{title_text}  範圍：{data.get('scope', '')}")
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("班級：___________  座號：_____  姓名：_______________  得分：___________")
    run.font.name = '微軟正黑體'
    run.font.size = Pt(11)
    
    # 一、選擇題
    mc_questions = data.get('mc_questions', [])
    if mc_questions:
        add_heading_2(doc, f"一、選擇題（{data.get('mc_scoring', '共 20 分')}）")
        for q in mc_questions:
            add_question(doc, q['question'], q['number'])
            
            if q.get('geometry'):
                insert_geometry_image(doc, f"mc_{q['number']}_fig", q, geom_dir)
                
            opts = q.get('options', {})
            p_opt = doc.add_paragraph()
            p_opt.paragraph_format.left_indent = Cm(1.0)
            
            p_opt.add_run("  ➀ ")
            add_math_runs(p_opt, opts.get('A',''))
            p_opt.add_run("   ➁ ")
            add_math_runs(p_opt, opts.get('B',''))
            p_opt.add_run("   ➂ ")
            add_math_runs(p_opt, opts.get('C',''))
            p_opt.add_run("   ➃ ")
            add_math_runs(p_opt, opts.get('D',''))
            
    # 二、填充題與計算題分類
    open_questions = data.get('open_questions', [])
    half = len(open_questions) // 2
    fill_questions = open_questions[:10] if len(open_questions) >= 10 else open_questions[:half]
    app_questions = open_questions[len(fill_questions):]
        
    if fill_questions:
        add_heading_2(doc, "二、填充題（共 40 分，每格 4 分）")
        for i, q in enumerate(fill_questions):
            num = i + 1
            ctx = q.get('context', '')
            p_q = doc.add_paragraph()
            p_q.paragraph_format.left_indent = Cm(0.5)
            p_q.add_run(f"{num}. ")
            add_math_runs(p_q, ctx)
            if q.get('sub_questions'):
                p_q.add_run(" ")
                add_math_runs(p_q, q['sub_questions'][0].get('question', ''))
            
            if q.get('geometry'):
                insert_geometry_image(doc, f"open_{q['number']}_fig", q, geom_dir)
            
    if app_questions:
        add_heading_2(doc, "三、計算與應用題（共 40 分，每題 5 分）")
        for i, q in enumerate(app_questions):
            num = i + 1
            ctx = q.get('context', '')
            p_q = doc.add_paragraph()
            p_q.paragraph_format.left_indent = Cm(0.5)
            p_q.add_run(f"{num}. ")
            add_math_runs(p_q, ctx)
            
            if q.get('geometry'):
                insert_geometry_image(doc, f"open_{q['number']}_fig", q, geom_dir)
                
            if len(q.get('sub_questions', [])) == 1:
                sq = q['sub_questions'][0]
                p_sub = doc.add_paragraph()
                p_sub.paragraph_format.left_indent = Cm(0.5)
                add_math_runs(p_sub, sq.get('question',''))
                
                if sq.get('geometry'):
                    insert_geometry_image(doc, f"open_{q['number']}_{sq.get('label','').strip('()')}_fig", sq, geom_dir)
            else:
                for sq in q.get('sub_questions', []):
                    p_sub = doc.add_paragraph()
                    p_sub.paragraph_format.left_indent = Cm(1.0)
                    p_sub.add_run(f"{sq.get('label','')} ")
                    add_math_runs(p_sub, sq.get('question',''))
                    
                    if sq.get('geometry'):
                        insert_geometry_image(doc, f"open_{q['number']}_{sq.get('label','').strip('()')}_fig", sq, geom_dir)
                    
            p_space = doc.add_paragraph("\n\n")
            
    en = data.get('exam_number', '2')
    grade = data.get('grade', '五年級')
    docx_path = out_dir / f"{grade}數學第{en}次段考_題目卷.docx"
    doc.save(docx_path)
    print(f"✅ 題目卷已生成: {docx_path}")

def create_answer_paper(data, out_dir):
    """產生答案卷（教師版）"""
    doc = Document()
    en = data.get('exam_number', '2')
    grade = data.get('grade', '五年級')
    
    add_title(doc, f"{grade} 數學科 答案卷（教師版）")
    add_subtitle(doc, f"範圍：{data.get('scope', '')}")
    
    # 選擇題答案
    mc_questions = data.get('mc_questions', [])
    if mc_questions:
        add_heading_2(doc, "一、選擇題答案")
        ans_map = {'A': '➀', 'B': '➁', 'C': '➂', 'D': '➃'}
        ans_list = [f"({q['number']}) {ans_map.get(q['answer'], q['answer'])}" for q in mc_questions]
        p = doc.add_paragraph("  ,  ".join(ans_list))
        p.paragraph_format.left_indent = Cm(0.5)
        for q in mc_questions:
            p_sol = doc.add_paragraph()
            p_sol.paragraph_format.left_indent = Cm(0.5)
            p_sol.add_run(f"第 {q['number']} 題解析：")
            add_math_runs(p_sol, q.get('solution',''))
            
    # 非選擇題答案
    open_questions = data.get('open_questions', [])
    if open_questions:
        add_heading_2(doc, "二、非選擇題與應用題答案與詳解")
        for q in open_questions:
            p_q = doc.add_paragraph()
            p_q.paragraph_format.left_indent = Cm(0.5)
            p_q.add_run(f"題號 {q.get('number')}：")
            add_math_runs(p_q, q.get('context',''))
            p_q.runs[0].bold = True
            
            if len(q.get('sub_questions', [])) == 1:
                sq = q['sub_questions'][0]
                p_ans = doc.add_paragraph()
                p_ans.paragraph_format.left_indent = Cm(0.5)
                add_math_runs(p_ans, sq.get('question',''))
                p_ans.add_run(" ➔ 答：")
                add_math_runs(p_ans, sq.get('answer',''))
                
                p_sol = doc.add_paragraph()
                p_sol.paragraph_format.left_indent = Cm(0.7)
                p_sol.add_run("解析：")
                add_math_runs(p_sol, sq.get('solution',''))
            else:
                for sq in q.get('sub_questions', []):
                    p_ans = doc.add_paragraph()
                    p_ans.paragraph_format.left_indent = Cm(1.0)
                    p_ans.add_run(f"{sq.get('label','')} ")
                    add_math_runs(p_ans, sq.get('question',''))
                    p_ans.add_run(" ➔ 答：")
                    add_math_runs(p_ans, sq.get('answer',''))
                    
                    p_sol = doc.add_paragraph()
                    p_sol.paragraph_format.left_indent = Cm(1.2)
                    p_sol.add_run("解析：")
                    add_math_runs(p_sol, sq.get('solution',''))
                
    docx_path = out_dir / f"{grade}數學第{en}次段考_答案卷（教師版）.docx"
    doc.save(docx_path)
    print(f"✅ 答案卷已生成: {docx_path}")

def create_blueprint_table(data, out_dir):
    """產生雙向細目表 Word"""
    doc = Document()
    en = data.get('exam_number', '2')
    grade = data.get('grade', '五年級')
    
    add_title(doc, f"{grade} 數學科 雙向細目表")
    add_subtitle(doc, f"範圍：{data.get('scope', '')}")
    
    table = doc.add_table(rows=1, cols=6)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.style = 'Table Grid'
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = '教材章節'
    hdr_cells[1].text = '記憶'
    hdr_cells[2].text = '理解'
    hdr_cells[3].text = '應用'
    hdr_cells[4].text = '分析'
    hdr_cells[5].text = '合計分值(比例)'
    
    sections = {}
    
    all_qs = []
    for q in data.get('mc_questions', []):
        all_qs.append({'section': q.get('section', '其他'), 'bloom': q.get('bloom_level', '第2級（理解）'), 'points': 4})
    
    open_qs = data.get('open_questions', [])
    half = len(open_qs) // 2
    for i, q in enumerate(open_qs):
        sec = q.get('section', '其他')
        is_fill = (i < 10) if len(open_qs) >= 10 else (i < half)
        pts = 4 if is_fill else 5
        for sq in q.get('sub_questions', []):
            all_qs.append({'section': sec, 'bloom': sq.get('bloom_level', q.get('bloom_level', '第3級（應用）')), 'points': sq.get('points', pts)})
            
    for q in all_qs:
        sec = q['section']
        bloom = q['bloom']
        pts = q['points']
        if sec not in sections:
            sections[sec] = {'記憶': 0, '理解': 0, '應用': 0, '分析': 0, 'total': 0}
        
        level = '理解'
        if '1' in bloom or '記憶' in bloom: level = '記憶'
        elif '3' in bloom or '應用' in bloom: level = '應用'
        elif '4' in bloom or '分析' in bloom: level = '分析'
        
        sections[sec][level] += pts
        sections[sec]['total'] += pts
        
    for sec, stats in sections.items():
        row_cells = table.add_row().cells
        row_cells[0].text = sec
        row_cells[1].text = str(stats['記憶']) if stats['記憶'] > 0 else '-'
        row_cells[2].text = str(stats['理解']) if stats['理解'] > 0 else '-'
        row_cells[3].text = str(stats['應用']) if stats['應用'] > 0 else '-'
        row_cells[4].text = str(stats['分析']) if stats['分析'] > 0 else '-'
        row_cells[5].text = f"{stats['total']} 分 ({stats['total']}%)"
        
    row_cells = table.add_row().cells
    row_cells[0].text = '合計分值'
    sum_m = sum(s['記憶'] for s in sections.values())
    sum_u = sum(s['理解'] for s in sections.values())
    sum_a = sum(s['應用'] for s in sections.values())
    sum_an = sum(s['分析'] for s in sections.values())
    
    row_cells[1].text = f"{sum_m} 分 ({sum_m}%)"
    row_cells[2].text = f"{sum_u} 分 ({sum_u}%)"
    row_cells[3].text = f"{sum_a} 分 ({sum_a}%)"
    row_cells[4].text = f"{sum_an} 分 ({sum_an}%)"
    row_cells[5].text = "100 分 (100%)"
    
    docx_path = out_dir / f"{grade}數學第{en}次段考_雙向細目表.docx"
    doc.save(docx_path)
    print(f"✅ 雙向細目表已生成: {docx_path}")

def main():
    if len(sys.argv) < 3:
        print("Usage: python generate_exam_docx.py <exam_data.json> <output_dir/> [--geom-dir <geom_dir>] [--blueprint-only]")
        sys.exit(1)
        
    json_path = Path(sys.argv[1])
    out_dir = Path(sys.argv[2])
    out_dir.mkdir(parents=True, exist_ok=True)
    
    bp_only = "--blueprint-only" in sys.argv
    
    geom_dir = None
    if "--geom-dir" in sys.argv:
        idx = sys.argv.index("--geom-dir")
        if idx + 1 < len(sys.argv):
            geom_dir = Path(sys.argv[idx + 1])
            
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
        
    create_blueprint_table(data, out_dir)
    if not bp_only:
        create_exam_paper(data, out_dir, geom_dir)
        create_answer_paper(data, out_dir)
        
if __name__ == '__main__':
    main()
